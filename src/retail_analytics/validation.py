from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import EXPECTED_COLUMNS, EXPECTED_SHEETS, RAW_WORKBOOK_PATH


class DatasetValidationError(ValueError):
    """Raised when the raw workbook does not match the expected source contract."""


@dataclass(frozen=True)
class SheetSummary:
    name: str
    data_rows: int


@dataclass(frozen=True)
class WorkbookSummary:
    path: Path
    sheets: tuple[SheetSummary, ...]

    @property
    def total_rows(self) -> int:
        return sum(sheet.data_rows for sheet in self.sheets)


def validate_workbook(path: Path = RAW_WORKBOOK_PATH) -> WorkbookSummary:
    """Validate workbook sheets and source columns without changing raw data."""
    path = Path(path)

    if not path.is_file():
        raise DatasetValidationError(f"Dataset workbook not found: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise DatasetValidationError(
            f"Dataset workbook could not be opened: {path}"
        ) from exc

    try:
        missing_sheets = [
            sheet_name
            for sheet_name in EXPECTED_SHEETS
            if sheet_name not in workbook.sheetnames
        ]
        if missing_sheets:
            missing = ", ".join(missing_sheets)
            raise DatasetValidationError(f"Missing required sheet(s): {missing}")

        summaries: list[SheetSummary] = []

        for sheet_name in EXPECTED_SHEETS:
            worksheet = workbook[sheet_name]
            first_row = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                None,
            )

            if first_row is None:
                raise DatasetValidationError(
                    f"Sheet {sheet_name!r} does not contain a header row"
                )

            actual_columns = tuple(first_row)
            if actual_columns != EXPECTED_COLUMNS:
                raise DatasetValidationError(
                    f"Unexpected columns in sheet {sheet_name!r}. "
                    f"Expected {EXPECTED_COLUMNS!r}, found {actual_columns!r}"
                )

            data_rows = max((worksheet.max_row or 1) - 1, 0)
            if data_rows == 0:
                raise DatasetValidationError(
                    f"Sheet {sheet_name!r} does not contain any data rows"
                )

            summaries.append(SheetSummary(name=sheet_name, data_rows=data_rows))

        return WorkbookSummary(path=path, sheets=tuple(summaries))
    finally:
        workbook.close()
